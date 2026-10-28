"""
OpenValora - I/O Module
Leitor e escritor de arquivos de avaliação em formato de planilha binária BIFF8 (.sda),
além de importação/exportação Excel (.xlsx/.xls) e CSV.
"""

import io
import re
from pathlib import Path
from typing import Tuple, Dict, Any, List, Optional
import xlrd
import xlwt
import openpyxl
import csv

from .models import (
    Project, ProjectMetadata, Variable, SampleItem,
    SubjectProperty, CalculationOptions
)
from .statistics import safe_float

def read_sda_file(file_path_or_bytes) -> Project:
    """Lê um arquivo .sda (BIFF8 Excel) e retorna um objeto Project."""
    if isinstance(file_path_or_bytes, (str, Path)):
        wb = xlrd.open_workbook(str(file_path_or_bytes))
    else:
        wb = xlrd.open_workbook(file_contents=file_path_or_bytes)

    project = Project()

    # 1. Sheet Version
    if "Version" in wb.sheet_names():
        sh_ver = wb.sheet_by_name("Version")
        ver_dict = {}
        for r in range(sh_ver.nrows):
            k = str(sh_ver.cell_value(r, 0)).strip()
            v = sh_ver.cell_value(r, 1)
            ver_dict[k] = v

        project.metadata.nome = str(ver_dict.get("Nome do Modelo", "Modelo OpenValora")).strip()
        raw_autor = str(ver_dict.get("Autor Modelo", "")).strip()
        if not raw_autor or any(w in raw_autor.lower() for w in ["sisdat", "responsavel", "responsável", "modelo", "exemplo"]):
            project.metadata.autor = "Nome do Responsável"
        else:
            project.metadata.autor = raw_autor
        project.metadata.data_criacao = str(ver_dict.get("Data Criação", "")).strip()
        raw_desc = str(ver_dict.get("Descrição do Modelo", "")).strip()
        project.metadata.descricao = raw_desc
        try:
            project.metadata.area_conhecimento = int(float(ver_dict.get("Área de Concentração", 1)))
        except (ValueError, TypeError):
            project.metadata.area_conhecimento = 1
        try:
            project.metadata.tipologia = int(float(ver_dict.get("Tipologia", 1)))
        except (ValueError, TypeError):
            project.metadata.tipologia = 1

    # 2. Sheet Var
    var_list: List[Variable] = []
    seen_var_names = set()
    if "Var" in wb.sheet_names():
        sh_var = wb.sheet_by_name("Var")
        # Header: Nome, Descrição, Tipo, Crescimento, Classificação, Continua, Decimais, Habilitada, Unitário, VarVinculada, ...
        for r in range(1, sh_var.nrows):
            nome = str(sh_var.cell_value(r, 0)).strip()
            if not nome or nome in seen_var_names:
                continue
            seen_var_names.add(nome)
            desc = str(sh_var.cell_value(r, 1)).strip() if sh_var.ncols > 1 else ""
            tipo = int(float(sh_var.cell_value(r, 2))) if sh_var.ncols > 2 and sh_var.cell_value(r, 2) != "" else 1
            cresc = int(float(sh_var.cell_value(r, 3))) if sh_var.ncols > 3 and sh_var.cell_value(r, 3) != "" else 0
            classif = int(float(sh_var.cell_value(r, 4))) if sh_var.ncols > 4 and sh_var.cell_value(r, 4) != "" else 1
            continua = int(float(sh_var.cell_value(r, 5))) if sh_var.ncols > 5 and sh_var.cell_value(r, 5) != "" else 0
            decimais = int(float(sh_var.cell_value(r, 6))) if sh_var.ncols > 6 and sh_var.cell_value(r, 6) != "" else 2
            hab = bool(float(sh_var.cell_value(r, 7))) if sh_var.ncols > 7 and sh_var.cell_value(r, 7) != "" else True
            unit = bool(float(sh_var.cell_value(r, 8))) if sh_var.ncols > 8 and sh_var.cell_value(r, 8) != "" else False
            var_vinc = int(float(sh_var.cell_value(r, 9))) if sh_var.ncols > 9 and sh_var.cell_value(r, 9) != "" else -1

            var_obj = Variable(
                nome=nome,
                descricao=desc,
                tipo=tipo,
                crescimento=cresc,
                classificacao=classif,
                continua=continua,
                decimais=decimais,
                habilitada=hab,
                unitario=unit,
                var_vinculada=var_vinc
            )
            var_list.append(var_obj)

    project.variaveis = var_list

    # 3. Sheet Options
    if "Options" in wb.sheet_names():
        sh_opt = wb.sheet_by_name("Options")
        opt_dict = {}
        for r in range(sh_opt.nrows):
            k = str(sh_opt.cell_value(r, 0)).strip()
            v = sh_opt.cell_value(r, 1) if sh_opt.ncols > 1 else None
            opt_dict[k.lower()] = v

        metodo_val = opt_dict.get("metodo", 1.0)
        project.opcoes.metodo = "simplificado" if metodo_val == 2.0 else "todas"
        
        sele_val = opt_dict.get("selemod", 1.0)
        project.opcoes.selecao_modelo = "regressao" if sele_val == 2.0 else "estimativa"
        
        mq_val = opt_dict.get("mq_mv", 1.0)
        project.opcoes.estimador = "mv" if mq_val == 2.0 else "mq"

        transf = []
        if opt_dict.get("x", 1.0) == 1.0: transf.append("x")
        if opt_dict.get("1/x", 0.0) == 1.0: transf.append("1/x")
        if opt_dict.get("ln(x)", 0.0) == 1.0: transf.append("ln(x)")
        if opt_dict.get("x2", 0.0) == 1.0: transf.append("x2")
        if opt_dict.get("1/x2", 0.0) == 1.0: transf.append("1/x2")
        if opt_dict.get("x(1/2)", 0.0) == 1.0: transf.append("x(1/2)")
        if opt_dict.get("1/x(1/2)", 0.0) == 1.0: transf.append("1/x(1/2)")
        if transf:
            project.opcoes.transformacoes = transf

    # 4. Sheet Data (Amostra)
    dados_list: List[SampleItem] = []
    if "Data" in wb.sheet_names():
        sh_data = wb.sheet_by_name("Data")
        for r in range(sh_data.nrows):
            hab_val = sh_data.cell_value(r, 0)
            habilitado = bool(float(hab_val)) if hab_val != "" else True
            out_val = sh_data.cell_value(r, 1) if sh_data.ncols > 1 else 0
            outlier = bool(float(out_val)) if out_val != "" else False

            valores = {}
            for col_idx, var in enumerate(project.variaveis):
                target_c = 2 + col_idx
                if target_c < sh_data.ncols:
                    cell_val = sh_data.cell_value(r, target_c)
                    if var.tipo == 1:  # Numérica
                        try:
                            valores[var.nome] = float(cell_val) if cell_val != "" else 0.0
                        except (ValueError, TypeError):
                            valores[var.nome] = 0.0
                    else:
                        valores[var.nome] = str(cell_val)
                else:
                    valores[var.nome] = 0.0 if var.tipo == 1 else ""

            dados_list.append(SampleItem(
                id=r + 1,
                habilitado=habilitado,
                outlier=outlier,
                valores=valores
            ))

    project.dados = dados_list

    # 5. Sheet Proj (Projeção / Imóveis Avaliandos)
    proj_list: List[SubjectProperty] = []
    if "Proj" in wb.sheet_names():
        sh_proj = wb.sheet_by_name("Proj")
        for r in range(sh_proj.nrows):
            valores = {}
            # As colunas em Proj mapeiam diretamente para as variáveis
            for col_idx, var in enumerate(project.variaveis):
                if col_idx < sh_proj.ncols:
                    cell_val = sh_proj.cell_value(r, col_idx)
                    if var.tipo == 1:
                        try:
                            valores[var.nome] = float(cell_val) if cell_val != "" else 0.0
                        except (ValueError, TypeError):
                            valores[var.nome] = 0.0
                    else:
                        valores[var.nome] = str(cell_val)
                else:
                    valores[var.nome] = 0.0 if var.tipo == 1 else ""

            proj_list.append(SubjectProperty(
                id=f"proj_{r+1}",
                identificacao=f"Imóvel Avaliando {r+1}",
                valores=valores
            ))

    project.projecoes = proj_list

    return project

def write_sda_file(project: Project) -> bytes:
    """Gera o arquivo binário .sda (BIFF8 Excel) compatível com sistemas de avaliação imobiliária."""
    wb = xlwt.Workbook(encoding="utf-8")

    # 1. Sheet Version
    sh_ver = wb.add_sheet("Version")
    ver_data = [
        ("Application", "OpenValora"),
        ("Packge", "Home"),
        ("Version", 1.0),
        ("Subversion", 0.1),
        ("Copyright", "OpenValora"),
        ("Autor Modelo", project.metadata.autor or "Nome do Responsável"),
        ("Data Criação", project.metadata.data_criacao),
        ("Senha", project.metadata.senha),
        ("Tipo Senha", float(project.metadata.tipo_senha)),
        ("Nome do Modelo", project.metadata.nome),
        ("Descrição do Modelo", project.metadata.descricao),
        ("Área de Concentração", float(project.metadata.area_conhecimento)),
        ("Tipologia", float(project.metadata.tipologia))
    ]
    for r, (k, v) in enumerate(ver_data):
        sh_ver.write(r, 0, k)
        sh_ver.write(r, 1, v)

    # 2. Sheet Page
    sh_page = wb.add_sheet("Page")
    page_data = [
        ("Cabeçalho", "OpenValora - Modelagem de Dados&b&b&d\n&w\n"),
        ("Rodapé", "\n&b&p/&P"),
        ("m_Esq", 20.0),
        ("m_Dir", 15.0),
        ("m_Sup", 10.0),
        ("m_Inf", 10.0),
        ("m_Cab", 20.0),
        ("m_Rod", 20.0),
        ("Data Impressão", 0.0),
        ("Número da Página", 0.0)
    ]
    for r, (k, v) in enumerate(page_data):
        sh_page.write(r, 0, k)
        sh_page.write(r, 1, v)

    # 3. Sheet Var
    sh_var = wb.add_sheet("Var")
    headers_var = [
        "Nome", "Descrição", "Tipo", "Crescimento", "Classificação",
        "Continua", "Decimais", "Habilitada", "Unitário", "VarVinculada",
        "LarguraColuna", "Transformacao RL"
    ]
    for c, h in enumerate(headers_var):
        sh_var.write(0, c, h)

    for r, var in enumerate(project.variaveis, start=1):
        sh_var.write(r, 0, var.nome)
        sh_var.write(r, 1, var.descricao)
        sh_var.write(r, 2, float(var.tipo))
        sh_var.write(r, 3, float(var.crescimento))
        sh_var.write(r, 4, float(var.classificacao))
        sh_var.write(r, 5, float(var.continua))
        sh_var.write(r, 6, float(var.decimais))
        sh_var.write(r, 7, 1.0 if var.habilitada else 0.0)
        sh_var.write(r, 8, 1.0 if var.unitario else 0.0)
        sh_var.write(r, 9, float(var.var_vinculada))
        sh_var.write(r, 10, float(var.largura_coluna))
        sh_var.write(r, 11, float(var.transformacao_rl))

    # 4. Sheet Model
    sh_model = wb.add_sheet("Model")
    nd = float(len(project.dados))
    nd_consid = float(len([d for d in project.dados if d.habilitado]))
    nd_proj = float(len(project.projecoes))
    nv = float(len(project.variaveis))
    nv_num = float(len([v for v in project.variaveis if v.tipo == 1]))
    nv_consid = float(len([v for v in project.variaveis if v.habilitada and v.tipo == 1]))

    sh_model.write(0, 1, "RL")
    model_data = [
        ("Nd", nd),
        ("NdConsid", nd_consid),
        ("NdProj", nd_proj),
        ("NdValid", 0.0),
        ("Nv", nv),
        ("NvConsid", nv_consid),
        ("NvNum", nv_num),
        ("Ne", 100.0),
        ("NeSelec", float(project.modelo_selecionado_idx)),
        ("Calculado", 1.0 if project.modelos_calculados else 0.0),
        ("Var Dep", 0.0)
    ]
    for r, (k, v) in enumerate(model_data, start=1):
        sh_model.write(r, 0, k)
        sh_model.write(r, 1, v)

    # 5. Sheet Options
    sh_opt = wb.add_sheet("Options")
    sh_opt.write(0, 1, "RL")
    t_set = set(project.opcoes.transformacoes)
    options_data = [
        ("Var Dep", ""),
        ("Metodo", 2.0 if project.opcoes.metodo == "simplificado" else 1.0),
        ("SeleMod", 2.0 if project.opcoes.selecao_modelo == "regressao" else 1.0),
        ("MQ_MV", 2.0 if project.opcoes.estimador == "mv" else 1.0),
        ("x", 1.0 if "x" in t_set else 0.0),
        ("1/x", 1.0 if "1/x" in t_set else 0.0),
        ("Ln(x)", 1.0 if "ln(x)" in t_set else 0.0),
        ("x2", 1.0 if "x2" in t_set else 0.0),
        ("1/x2", 1.0 if "1/x2" in t_set else 0.0),
        ("x(1/2)", 1.0 if "x(1/2)" in t_set else 0.0),
        ("1/x(1/2)", 1.0 if "1/x(1/2)" in t_set else 0.0),
        ("Saida Linear", ""),
        ("Bagging", ""),
        ("Prune", "")
    ]
    for r, (k, v) in enumerate(options_data, start=1):
        sh_opt.write(r, 0, k)
        sh_opt.write(r, 1, v)

    # 6. Sheet RL
    sh_rl = wb.add_sheet("RL")
    sh_rl.write(0, 0, "EQ")

    # 7. Sheet RNA
    sh_rna = wb.add_sheet("RNA")
    sh_rna.write(0, 0, "nvEntrada")

    # 8. Sheet RNL
    sh_rnl = wb.add_sheet("RNL")
    sh_rnl.write(0, 0, "Iterações")

    # 9. Sheet DEA
    sh_dea = wb.add_sheet("DEA")
    sh_dea.write(0, 0, "Opção")
    sh_dea.write(0, 1, "R2")

    # 10. Sheet Data
    sh_data = wb.add_sheet("Data")
    for r, item in enumerate(project.dados):
        sh_data.write(r, 0, 1.0 if item.habilitado else 0.0)
        sh_data.write(r, 1, 1.0 if item.outlier else 0.0)
        for c, var in enumerate(project.variaveis):
            val = item.valores.get(var.nome, "" if var.tipo != 1 else 0.0)
            if var.tipo == 1:
                sh_data.write(r, 2 + c, safe_float(val))
            else:
                sh_data.write(r, 2 + c, str(val))

    # 11. Sheet Proj
    sh_proj = wb.add_sheet("Proj")
    for r, item in enumerate(project.projecoes):
        for c, var in enumerate(project.variaveis):
            val = item.valores.get(var.nome, "" if var.tipo != 1 else 0.0)
            if var.tipo == 1:
                sh_proj.write(r, c, safe_float(val))
            else:
                sh_proj.write(r, c, str(val))

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def decode_text_bytes(file_bytes: bytes) -> str:
    """Decodifica texto tentando UTF-8, CP1252 (Windows) e Latin-1."""
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return file_bytes.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return file_bytes.decode("utf-8", errors="replace")

def import_excel_or_csv(file_bytes: bytes, filename: str) -> Tuple[List[Variable], List[SampleItem]]:
    """Importa uma tabela arbitrária Excel ou CSV gerando variáveis e dados."""
    fn = filename.lower()
    rows_data = []

    if fn.endswith(".csv"):
        text = decode_text_bytes(file_bytes)
        reader = csv.reader(io.StringIO(text), delimiter=";" if ";" in text else ",")
        rows_data = list(reader)
    elif fn.endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sh = wb.active
        for row in sh.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                rows_data.append(list(row))
    elif fn.endswith(".xls") or fn.endswith(".sda"):
        wb = xlrd.open_workbook(file_contents=file_bytes)
        sh = wb.sheet_by_name("Data") if "Data" in wb.sheet_names() else wb.sheet_by_index(0)
        for r in range(sh.nrows):
            rows_data.append([sh.cell_value(r, c) for c in range(sh.ncols)])

    if not rows_data or len(rows_data) < 2:
        return [], []

    # Headers com garantia de unicidade
    clean_headers = []
    header_counts = {}
    for i, h in enumerate(rows_data[0]):
        raw_name = str(h).strip() if h is not None and str(h).strip() else f"Var_{i+1}"
        if raw_name in header_counts:
            header_counts[raw_name] += 1
            clean_headers.append(f"{raw_name}_{header_counts[raw_name]}")
        else:
            header_counts[raw_name] = 1
            clean_headers.append(raw_name)
    headers = clean_headers
    variables: List[Variable] = []
    
    # Inferência de tipos
    for col_idx, h in enumerate(headers):
        is_num = True
        numeric_count = 0
        total_non_empty = 0
        for r in range(1, len(rows_data)):
            val = rows_data[r][col_idx] if col_idx < len(rows_data[r]) else None
            if val is not None and str(val).strip() != "":
                total_non_empty += 1
                s_val = str(val).strip()
                # Se contém letras alfabéticas (exceto prefixos de moeda R$, etc), é texto
                cleaned = re.sub(r"[R$\s\.,%-]", "", s_val)
                if any(c.isalpha() for c in cleaned):
                    is_num = False
                    break
                try:
                    safe_float(s_val)
                    numeric_count += 1
                except Exception:
                    is_num = False
                    break
        
        is_numeric = is_num and (numeric_count > 0)
        var_type = 1 if is_numeric else 2
        # Se última coluna numérica, assumir dependente preliminar
        is_last = (col_idx == len(headers) - 1)
        classif = 7 if (is_numeric and is_last) else (1 if is_numeric else 8)

        var = Variable(
            nome=h,
            descricao=f"Variável importada: {h}",
            tipo=var_type,
            classificacao=classif,
            continua=1 if is_numeric else 0,
            decimais=2,
            habilitada=True,
            unitario=True if ("m2" in h.lower() or "m²" in h.lower() or "unit" in h.lower()) else False
        )
        variables.append(var)

    samples: List[SampleItem] = []
    for r in range(1, len(rows_data)):
        row_vals = rows_data[r]
        val_map = {}
        for col_idx, var in enumerate(variables):
            raw = row_vals[col_idx] if col_idx < len(row_vals) else ""
            if var.tipo == 1:
                val_map[var.nome] = safe_float(raw) if raw is not None and str(raw).strip() != "" else 0.0
            else:
                val_map[var.nome] = str(raw).strip() if raw is not None else ""

        samples.append(SampleItem(
            id=r,
            habilitado=True,
            outlier=False,
            valores=val_map
        ))

    return variables, samples

def export_to_excel(project: Project) -> bytes:
    """Exporta o projeto completo com dados e estatísticas para um arquivo Excel (.xlsx)."""
    wb = openpyxl.Workbook()
    # Sheet 1: Amostra
    ws_data = wb.active
    ws_data.title = "Dados da Amostra"
    
    headers = ["ID", "Habilitado", "Outlier"] + [v.nome for v in project.variaveis]
    ws_data.append(headers)

    for item in project.dados:
        row = [item.id, "Sim" if item.habilitado else "Não", "Sim" if item.outlier else "Não"]
        for v in project.variaveis:
            row.append(item.valores.get(v.nome, ""))
        ws_data.append(row)

    # Sheet 2: Variáveis
    ws_var = wb.create_sheet(title="Variáveis")
    ws_var.append(["Nome", "Descrição", "Tipo", "Classificação", "Crescimento", "Decimais", "Habilitada", "Unitário"])
    tipo_map = {1: "Numérica", 2: "Texto", 3: "Data", 4: "Moeda", 5: "Porcentagem"}
    classif_map = {1: "Quantitativa", 2: "Qualitativa", 3: "Proxy", 4: "Dicotômica", 7: "Dependente", 8: "Texto"}
    cresc_map = {0: "Crescente (Direta)", 1: "Decrescente (Inversa)", -1: "Livre"}

    for v in project.variaveis:
        ws_var.append([
            v.nome, v.descricao, tipo_map.get(v.tipo, "Numérica"),
            classif_map.get(v.classificacao, "Quantitativa"),
            cresc_map.get(v.crescimento, "Livre"),
            v.decimais, "Sim" if v.habilitada else "Não", "Sim" if v.unitario else "Não"
        ])

    # Sheet 3: Projeções / Imóvel Avaliando
    if project.projecoes:
        ws_proj = wb.create_sheet(title="Imóvel Avaliando")
        ws_proj.append(["Identificação"] + [v.nome for v in project.variaveis])
        for p in project.projecoes:
            ws_proj.append([p.identificacao] + [p.valores.get(v.nome, "") for v in project.variaveis])

    # Sheet 4: Modelo de Regressão Selecionado (se houver)
    if project.modelos_calculados:
        idx = min(project.modelo_selecionado_idx, len(project.modelos_calculados) - 1)
        m = project.modelos_calculados[idx]
        ws_mod = wb.create_sheet(title="Modelo de Regressão")
        ws_mod.append(["Equação", m.equacao_texto])
        ws_mod.append(["R", m.r, "R²", m.r2, "R² Ajustado", m.r2_ajustado, "R² Original", m.r2_original])
        ws_mod.append(["Erro Padrão", m.erro_padrao, "Durbin-Watson", m.durbin_watson, "Status DW", m.durbin_watson_status])
        ws_mod.append([])
        ws_mod.append(["Regressores (Teste t de Student)"])
        ws_mod.append(["Variável", "Transformação", "Coeficiente", "Erro Padrão", "t Calc", "Significância (%)", "Elasticidade", "VIF"])
        for reg in m.regressores:
            ws_mod.append([
                reg.nome_exibicao, reg.transformacao_codigo, reg.coeficiente,
                reg.erro_padrao, reg.t_calc, reg.significancia_pct,
                reg.elasticidade_media, reg.vif
            ])
        if m.anova:
            ws_mod.append([])
            ws_mod.append(["ANOVA (Análise de Variância)"])
            ws_mod.append(["Fonte de Variação", "GL", "Soma dos Quadrados (SQ)", "Quadrado Médio (QM)", "F Calculado", "Significância F (%)"])
            ws_mod.append(["Regressão", m.anova.df_reg, m.anova.ss_reg, m.anova.ms_reg, m.anova.f_calc, m.anova.significancia_f_pct])
            ws_mod.append(["Resíduos", m.anova.df_res, m.anova.ss_res, m.anova.ms_res, "-", "-"])
            ws_mod.append(["Total", m.anova.df_tot, m.anova.ss_tot, "-", "-", "-"])
        if m.nbr:
            ws_mod.append([])
            ws_mod.append(["Enquadramento NBR 14653-2"])
            ws_mod.append(["Grau de Fundamentação", m.nbr.grau_fundamentacao])
            ws_mod.append(["Grau de Precisão", m.nbr.grau_precisao])
            ws_mod.append(["Amplitude do IC a 80%", f"{m.nbr.amplitude_ic_pct}%"])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
